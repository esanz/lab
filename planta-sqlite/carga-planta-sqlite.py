#!/usr/bin/env python3
"""
carga-planta-sqlite.py
======================

Carga la hoja "Planta de datos" de la hoja de cálculo de Planta de Datos en
una base de datos SQLite, para poder consultarla con busca-planta-sqlite.py
(o con cualquier cliente SQLite).

Tablas que se crean (se borran y se vuelven a crear en cada carga):

  planta   Una fila por cada fila no vacía de la hoja, con TODAS sus columnas
           (como TEXT, con el mismo nombre que en la hoja) más:
             fila  INTEGER PRIMARY KEY  -> número de fila en la hoja Excel
           Si una cabecera está repetida en la hoja, la segunda aparición se
           guarda como "<cabecera> (2)".

  rangos   Las redes CIDR de la columna "Rango Direccionamiento Datos"
           (una fila por red; una celda puede tener varias separadas por ";"):
             fila    -> planta.fila
             cidr    -> texto de la red normalizada (p.ej. 10.13.2.0/24)
             prefijo -> longitud del prefijo (24)
             inicio  -> primera dirección de la red, como entero
             fin     -> última dirección de la red, como entero

  carga    Metadatos de la última carga (fichero, hoja, fecha, nº de filas).

Índices:
  rangos(prefijo, inicio)  -> la búsqueda IP -> CIDR se resuelve con una
                              consulta por igualdad para cada longitud de
                              prefijo presente (red de la IP con ese prefijo),
                              así que siempre va por índice.
  rangos(fila)             -> unión con planta.
  planta("Código de Punto de Servicio")  -> consultas por PS.

Uso:
    python3 carga-planta-sqlite.py Planta_Datos.xlsx
    python3 carga-planta-sqlite.py Planta_Datos.xlsx -d /ruta/planta.sqlite

Requiere: openpyxl (pip install openpyxl)
"""

import argparse
import datetime as dt
import ipaddress
import os
import re
import sqlite3
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Falta el módulo openpyxl. Instálalo con: pip install openpyxl")


HOJA_POR_DEFECTO = "Planta de datos"
COLUMNA_RANGO = "Rango Direccionamiento Datos"
COLUMNA_PS = "Código de Punto de Servicio"
BD_POR_DEFECTO = "planta.sqlite"


def texto(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def q(nombre):
    """Identificador SQL entre comillas dobles."""
    return '"' + nombre.replace('"', '""') + '"'


def nombres_columnas(cabecera):
    nombres, usados = [], set()
    for i, c in enumerate(cabecera, start=1):
        base = texto(c) or f"Columna {i}"
        nombre, n = base, 1
        while nombre.lower() in usados or nombre.lower() == "fila":
            n += 1
            nombre = f"{base} ({n})"
        usados.add(nombre.lower())
        nombres.append(nombre)
    return nombres


def redes_de(valor, num_fila):
    redes = []
    for s in re.split(r"[;,\s]+", valor or ""):
        if not s:
            continue
        try:
            red = ipaddress.ip_network(s, strict=False)
        except ValueError:
            print(f"Aviso: fila {num_fila}: '{s}' no es un rango CIDR válido; se ignora.",
                  file=sys.stderr)
            continue
        if red.version != 4:
            print(f"Aviso: fila {num_fila}: '{s}' no es IPv4; se ignora.", file=sys.stderr)
            continue
        if red.prefixlen > 0:  # 0.0.0.0/0 es "vacío" en la hoja
            redes.append(red)
    return redes


def main():
    ap = argparse.ArgumentParser(
        prog="carga-planta-sqlite",
        description="Carga la Planta de Datos (.xlsx) en una base de datos SQLite.")
    ap.add_argument("planta", help="hoja de cálculo de la Planta de Datos (.xlsx)")
    ap.add_argument("-d", "--bd", default=BD_POR_DEFECTO,
                    help=f"base de datos SQLite de destino (por defecto {BD_POR_DEFECTO})")
    ap.add_argument("-H", "--hoja", help=f"nombre de la hoja (por defecto '{HOJA_POR_DEFECTO}')")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.planta, read_only=True, data_only=True)
    if args.hoja:
        if args.hoja not in wb.sheetnames:
            sys.exit(f"No existe la hoja '{args.hoja}'. Hojas: {', '.join(wb.sheetnames)}")
        ws = wb[args.hoja]
    elif HOJA_POR_DEFECTO in wb.sheetnames:
        ws = wb[HOJA_POR_DEFECTO]
    else:
        ws = wb.worksheets[0]

    filas = ws.iter_rows(values_only=True)
    columnas = nombres_columnas(next(filas, []))
    for requerida in (COLUMNA_RANGO, COLUMNA_PS):
        if requerida not in columnas:
            sys.exit(f"En la hoja '{ws.title}' falta la columna '{requerida}'")
    i_rango = columnas.index(COLUMNA_RANGO)

    con = sqlite3.connect(args.bd)
    try:
        with con:
            con.executescript("""
                DROP TABLE IF EXISTS rangos;
                DROP TABLE IF EXISTS planta;
                DROP TABLE IF EXISTS carga;
            """)
            con.execute("CREATE TABLE planta (fila INTEGER PRIMARY KEY, "
                        + ", ".join(f"{q(c)} TEXT" for c in columnas) + ")")
            con.execute("""
                CREATE TABLE rangos (
                    fila    INTEGER NOT NULL REFERENCES planta(fila),
                    cidr    TEXT    NOT NULL,
                    prefijo INTEGER NOT NULL,
                    inicio  INTEGER NOT NULL,
                    fin     INTEGER NOT NULL
                )""")
            con.execute("""
                CREATE TABLE carga (
                    fichero TEXT, hoja TEXT, fecha TEXT, filas INTEGER, rangos INTEGER
                )""")

            sql_planta = (f"INSERT INTO planta (fila, {', '.join(q(c) for c in columnas)}) "
                          f"VALUES ({', '.join('?' * (len(columnas) + 1))})")
            sql_rango = "INSERT INTO rangos VALUES (?, ?, ?, ?, ?)"

            n_filas = n_rangos = 0
            for num_fila, fila in enumerate(filas, start=2):
                valores = [texto(v) for v in fila][:len(columnas)]
                valores += [None] * (len(columnas) - len(valores))
                if all(v is None for v in valores):
                    continue
                con.execute(sql_planta, [num_fila] + valores)
                n_filas += 1
                for red in redes_de(valores[i_rango], num_fila):
                    con.execute(sql_rango, (num_fila, str(red), red.prefixlen,
                                            int(red.network_address),
                                            int(red.broadcast_address)))
                    n_rangos += 1

            con.executescript(f"""
                CREATE INDEX ix_rangos_prefijo_inicio ON rangos(prefijo, inicio);
                CREATE INDEX ix_rangos_fila ON rangos(fila);
                CREATE INDEX ix_planta_ps ON planta({q(COLUMNA_PS)});
            """)
            con.execute("INSERT INTO carga VALUES (?, ?, ?, ?, ?)",
                        (os.path.abspath(args.planta), ws.title,
                         dt.datetime.now().isoformat(sep=" ", timespec="seconds"),
                         n_filas, n_rangos))
        con.execute("ANALYZE")
    finally:
        con.close()
        wb.close()

    print(f"Cargadas {n_filas} filas ({len(columnas)} columnas) y {n_rangos} rangos CIDR "
          f"de la hoja '{ws.title}' en {args.bd}")


if __name__ == "__main__":
    main()
