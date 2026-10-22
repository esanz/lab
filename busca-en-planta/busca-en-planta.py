#!/usr/bin/env python3
"""
busca-en-planta.py
==================

Busca una o varias direcciones IP en la hoja de cálculo de "Planta de Datos"
y muestra, para cada IP, el/los Punto(s) de Servicio (PS) cuya red de datos
la contiene:

    <IP> --> <Código PS>, <Nombre del PS>, <Dirección>, <Municipio>,
             <Provincia>, <Código Postal>, <Coubicado>

Además del listado por pantalla genera un fichero CSV con el resultado.

Cómo se busca
-------------
La IP se compara con la columna "Rango Direccionamiento Datos", que contiene
redes en formato CIDR (p.ej. 10.13.2.0/24; puede haber varias separadas por
";"). Una fila encaja si la IP pertenece a alguna de sus redes, y se listan
TODAS las filas que encajen, cada una en su línea. Se ignora 0.0.0.0/0, que
la hoja usa como "vacío".

Uso:
    python3 busca-en-planta.py -p Planta_Datos.xlsx 10.13.2.2 10.238.80.77
    python3 busca-en-planta.py -p Planta_Datos.xlsx -f ips.txt
    python3 busca-en-planta.py -p Planta_Datos.xlsx ips.txt -o salida.csv
    cat ips.txt | python3 busca-en-planta.py -p Planta_Datos.xlsx -

Los argumentos posicionales pueden ser IPs o ficheros; de un fichero se
extraen todas las IPs que aparezcan (una por línea, separadas por comas,
con comentarios "#", etc.).

Requiere: openpyxl (pip install openpyxl)
"""

import argparse
import csv
import ipaddress
import os
import re
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    sys.exit("Falta el módulo openpyxl. Instálalo con: pip install openpyxl")


HOJA_POR_DEFECTO = "Planta de datos"
COLUMNA_RANGO = "Rango Direccionamiento Datos"

# (clave interna, cabecera en la hoja, cabecera en el CSV de salida)
CAMPOS = [
    ("codigo", "Código de Punto de Servicio", "Código de Punto de Servicio"),
    ("nombre", "Nombre del PS", "Nombre del PS"),
    ("direccion", "Dirección", "Dirección"),
    ("municipio", "Municipio", "Municipio"),
    ("provincia", "Provincia", "Provincia"),
    ("cp", "Código Postal", "Código Postal"),
    ("coubicado", "Coubicado", "Coubicado"),
]

RE_IP = re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")


def texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ---------------------------------------------------------------------------
# Lectura de la planta
# ---------------------------------------------------------------------------

def cargar_planta(ruta, nombre_hoja=None):
    """Devuelve una lista de (registro, redes, fila_excel).

    registro: dict con los CAMPOS de la fila.
    redes:    lista de ip_network de la columna COLUMNA_RANGO.
    """
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)

    if nombre_hoja:
        if nombre_hoja not in wb.sheetnames:
            sys.exit(f"No existe la hoja '{nombre_hoja}'. Hojas: {', '.join(wb.sheetnames)}")
        ws = wb[nombre_hoja]
    elif HOJA_POR_DEFECTO in wb.sheetnames:
        ws = wb[HOJA_POR_DEFECTO]
    else:
        ws = wb.worksheets[0]

    filas = ws.iter_rows(values_only=True)
    cabecera = [texto(c) for c in next(filas, [])]

    # Primera aparición de cada cabecera (hay alguna repetida en la hoja)
    pos = {}
    for i, c in enumerate(cabecera):
        pos.setdefault(c, i)

    faltan = [cab for cab in [COLUMNA_RANGO] + [cab for _, cab, _ in CAMPOS] if cab not in pos]
    if faltan:
        sys.exit(f"En la hoja '{ws.title}' faltan las columnas: {', '.join(faltan)}")

    planta = []
    for num_fila, fila in enumerate(filas, start=2):
        if not fila or all(v is None for v in fila):
            continue
        fila = list(fila) + [None] * (len(cabecera) - len(fila))

        redes = []
        for s in re.split(r"[;,\s]+", texto(fila[pos[COLUMNA_RANGO]])):
            if not s:
                continue
            try:
                red = ipaddress.ip_network(s, strict=False)
            except ValueError:
                print(f"Aviso: fila {num_fila}: '{s}' no es un rango CIDR válido; se ignora.",
                      file=sys.stderr)
                continue
            if red.prefixlen > 0:
                redes.append(red)
        if not redes:
            continue

        reg = {clave: texto(fila[pos[cab]]) for clave, cab, _ in CAMPOS}
        planta.append((reg, redes, num_fila))

    wb.close()
    return planta


# ---------------------------------------------------------------------------
# Búsqueda
# ---------------------------------------------------------------------------

def buscar(ip, planta):
    """Devuelve [(registro, [redes que contienen la IP], fila_excel), ...]."""
    resultado = []
    for reg, redes, num_fila in planta:
        encajan = [red for red in redes if ip in red]
        if encajan:
            resultado.append((reg, encajan, num_fila))
    return resultado


# ---------------------------------------------------------------------------
# Entrada de IPs
# ---------------------------------------------------------------------------

def ips_de_texto(contenido):
    ips = []
    for linea in contenido.splitlines():
        linea = linea.split("#", 1)[0]
        ips.extend(RE_IP.findall(linea))
    return ips


def recoger_ips(args):
    brutas = []
    for a in args.ips:
        if a == "-":
            brutas.extend(ips_de_texto(sys.stdin.read()))
        elif os.path.isfile(a):
            with open(a, encoding="utf-8-sig", errors="replace") as f:
                brutas.extend(ips_de_texto(f.read()))
        else:
            brutas.append(a)
    for fichero in args.fichero or []:
        with open(fichero, encoding="utf-8-sig", errors="replace") as f:
            brutas.extend(ips_de_texto(f.read()))

    ips, vistas = [], set()
    for s in brutas:
        s = s.strip()
        try:
            ip = ipaddress.ip_address(s)
        except ValueError:
            print(f"Aviso: '{s}' no es una dirección IP válida; se ignora.", file=sys.stderr)
            continue
        if ip not in vistas:
            vistas.add(ip)
            ips.append(ip)
    return ips


# ---------------------------------------------------------------------------
# Principal
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(
        prog="busca-en-planta",
        description="Busca direcciones IP en la Planta de Datos y muestra el Punto de Servicio.")
    ap.add_argument("-p", "--planta", required=True,
                    help="hoja de cálculo de la Planta de Datos (.xlsx)")
    ap.add_argument("ips", nargs="*",
                    help="direcciones IP, ficheros con IPs, o '-' para leer de la entrada estándar")
    ap.add_argument("-f", "--fichero", action="append",
                    help="fichero con direcciones IP (se puede repetir)")
    ap.add_argument("-H", "--hoja", help=f"nombre de la hoja (por defecto '{HOJA_POR_DEFECTO}')")
    ap.add_argument("-o", "--csv",
                    help="fichero CSV de salida (por defecto busca-en-planta_AAAAMMDD-HHMMSS.csv)")
    ap.add_argument("--sin-csv", action="store_true", help="no generar el fichero CSV")
    ap.add_argument("-v", "--detalle", action="store_true",
                    help="muestra en pantalla el rango CIDR y la fila de la hoja que encajan")
    args = ap.parse_args()

    ips = recoger_ips(args)
    if not ips:
        ap.error("no se ha indicado ninguna dirección IP válida")

    planta = cargar_planta(args.planta, args.hoja)

    filas_csv = []
    for ip in ips:
        res = buscar(ip, planta)
        if not res:
            print(f"{ip} --> NO ENCONTRADA")
            filas_csv.append([str(ip)] + [""] * len(CAMPOS) + ["", ""])
            continue
        for reg, redes, num_fila in res:
            valores = [reg[c] for c, _, _ in CAMPOS]
            rangos = "; ".join(str(r) for r in redes)
            linea = f"{ip} --> {', '.join(valores)}"
            if args.detalle:
                linea += f"  [{rangos}, fila {num_fila}]"
            print(linea)
            filas_csv.append([str(ip)] + valores + [rangos, num_fila])

    if not args.sin_csv:
        ruta = args.csv or datetime.now().strftime("busca-en-planta_%Y%m%d-%H%M%S.csv")
        # utf-8-sig y ';' para que Excel en español lo abra directamente
        with open(ruta, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Dirección IP"] + [c for _, _, c in CAMPOS]
                       + [COLUMNA_RANGO, "Fila en la hoja"])
            w.writerows(filas_csv)
        print(f"\nCSV generado: {ruta}", file=sys.stderr)


if __name__ == "__main__":
    main()
